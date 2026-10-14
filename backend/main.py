from datetime import datetime, timedelta
import base64
import csv
import hashlib
from io import BytesIO, StringIO
import os
from pathlib import Path
import re
import shutil
import secrets
import subprocess
import tempfile
import json
import mimetypes
from typing import Optional
from urllib.parse import urlparse
import unicodedata
import zipfile
from uuid import uuid4
from urllib.request import Request as UrlRequest, urlopen
from fastapi import FastAPI, Depends, HTTPException, status, Header, UploadFile, File, Form, Response, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from jose import jwt, JWTError
from passlib.context import CryptContext
from PIL import Image, UnidentifiedImageError
from fpdf import FPDF
from openpyxl import Workbook
import pyotp
import qrcode
from sqlalchemy import text
from sqlalchemy.orm import Session

import models
import database
from mailer import write_email
from schemas import (
    ErasmusCallCreate,
    ErasmusCallOut,
    ErasmusCallUpdate,
    AuthRequest,
    ChangePasswordRequest,
    PasswordCreateRequest,
    AuthResponse,
    UserOut,
    ApplicationCreate,
    ApplicationOut,
    ApplicationDocumentOut,
    ApplicationDetailOut,
    ApplicationReviewRequest,
    ApplicationStatusHistoryOut,
    NotificationOut,
    InboxMessageOut,
    InboxUnreadCountOut,
    InboxBulkDeleteRequest,
    InboxBulkActionOut,
    AuditLogOut,
    LoginAuditOut,
    BulkReviewRequest,
    BulkReviewResult,
    MessageOut,
    VerifyEmailRequest,
    ResendVerificationRequest,
    PasswordResetRequest,
    PasswordResetConfirm,
    TwoFactorCodeRequest,
    TwoFactorSetupOut,
    TwoFactorVerifyRequest,
    PushTokenRequest,
)

ENVIRONMENT = os.getenv("ENVIRONMENT", "development").strip().lower()
SECRET_KEY = os.getenv("SECRET_KEY", "dev-secret-change-me").strip()
ADMIN_EMAIL = os.getenv("ADMIN_EMAIL", "admin@ibu.edu.ba").strip().lower()
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "admin123")
cors_origins = os.getenv(
    "CORS_ORIGINS",
    "http://localhost:3000,http://127.0.0.1:3000",
)
CORS_ORIGINS = [origin.strip() for origin in cors_origins.split(",") if origin.strip()]

if ENVIRONMENT == "production":
    if SECRET_KEY == "dev-secret-change-me":
        raise RuntimeError("SECRET_KEY must be configured in production")
    if ADMIN_PASSWORD == "admin123":
        raise RuntimeError("ADMIN_PASSWORD must be changed in production")
    if not CORS_ORIGINS:
        raise RuntimeError("CORS_ORIGINS must contain the deployed web origin")

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

models.Base.metadata.create_all(bind=database.engine)
database.ensure_schema()

ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60 * 24
ALLOWED_EMAIL_DOMAINS = ("@ibu.edu.ba", "@stu.ibu.edu.ba")
WEB_APP_URL = os.getenv("WEB_APP_URL", "http://localhost:3000")
UPLOAD_DIR = Path(__file__).resolve().parent / "uploads"
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".jpg", ".jpeg", ".png"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png"}
ALLOWED_MIME_TYPES = {
    ".pdf": {"application/pdf"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    ".jpg": {"image/jpeg"},
    ".jpeg": {"image/jpeg"},
    ".png": {"image/png"},
}
FILE_SIGNATURES = {
    ".pdf": b"%PDF-",
    ".docx": b"PK\x03\x04",
    ".jpg": b"\xff\xd8\xff",
    ".jpeg": b"\xff\xd8\xff",
    ".png": b"\x89PNG\r\n\x1a\n",
}
DOC_TYPES = [
    "cv",
    "transcript",
    "motivation",
    "language_proof",
    "recommendation",
    "learning_agreement",
    "id_copy",
    "application_form",
]
DOC_LABELS = {
    "cv": "CV (Europass)",
    "transcript": "Transcript of Records",
    "motivation": "Motivation Letter",
    "language_proof": "Proof of Language",
    "recommendation": "Recommendation Letter",
    "learning_agreement": "Learning Agreement",
    "id_copy": "Passport/ID Copy",
    "application_form": "Application Form",
}
MAX_DOC_SIZE = 10 * 1024 * 1024
MAX_IMAGE_SIZE = 5 * 1024 * 1024
MAX_IMAGE_RAW_SIZE = 15 * 1024 * 1024
ENABLE_IMAGE_COMPRESSION = os.getenv("ENABLE_IMAGE_COMPRESSION", "true").lower() == "true"
CLAMAV_PATH = os.getenv("CLAMAV_PATH")
CLAMAV_TIMEOUT = int(os.getenv("CLAMAV_TIMEOUT", "20"))
NAME_PATTERN = re.compile(r"^[A-Za-zÀ-ÖØ-öø-ÿĀ-ž .'-]+$")
PHONE_PATTERN = re.compile(r"^\+?\d{6,20}$")
ALERT_DAYS = {7, 3, 1}
LOGIN_MAX_ATTEMPTS = 3
LOGIN_LOCK_MINUTES = 15
EMAIL_VERIFICATION_EXPIRY_HOURS = int(os.getenv("EMAIL_VERIFICATION_EXPIRY_HOURS", "24"))
PASSWORD_RESET_EXPIRY_MINUTES = int(os.getenv("PASSWORD_RESET_EXPIRY_MINUTES", "60"))

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")

UPLOAD_DIR.mkdir(exist_ok=True)

def get_db():
    db = database.SessionLocal()
    try:
        yield db
    finally:
        db.close()

def hash_password(password: str) -> str:
    return pwd_context.hash(password)

def verify_password(plain_password: str, password_hash: str) -> bool:
    return pwd_context.verify(plain_password, password_hash)

def create_access_token(subject: str) -> str:
    expire = datetime.utcnow() + timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    payload = {"sub": subject, "exp": expire}
    return jwt.encode(payload, SECRET_KEY, algorithm=ALGORITHM)

def create_two_factor_token(subject: str) -> str:
    expire = datetime.utcnow() + timedelta(minutes=5)
    payload = {"sub": subject, "purpose": "2fa", "exp": expire}
    return jwt.encode(payload, SECRET_KEY, algorithm=ALGORITHM)

def build_qr_code(data: str) -> str:
    image = qrcode.make(data)
    output = BytesIO()
    image.save(output, format="PNG")
    encoded = base64.b64encode(output.getvalue()).decode("ascii")
    return f"data:image/png;base64,{encoded}"

def create_token() -> str:
    return secrets.token_urlsafe(32)

def hash_token(token: str) -> str:
    return hashlib.sha256(token.encode("utf-8")).hexdigest()

def issue_email_verification(db: Session, user: models.User) -> None:
    db.query(models.EmailVerificationToken).filter(
        models.EmailVerificationToken.user_id == user.id,
        models.EmailVerificationToken.used_at.is_(None),
    ).delete()
    raw_token = create_token()
    token_hash = hash_token(raw_token)
    expires_at = datetime.utcnow() + timedelta(hours=EMAIL_VERIFICATION_EXPIRY_HOURS)
    db.add(
        models.EmailVerificationToken(
            user_id=user.id,
            token_hash=token_hash,
            expires_at=expires_at,
        )
    )
    db.commit()
    verify_link = f"{WEB_APP_URL}/verify?token={raw_token}"
    subject = "Verify your Erasmus account"
    body = (
        "Please verify your email to activate your Erasmus account.\n\n"
        f"{verify_link}\n\n"
        f"This link expires in {EMAIL_VERIFICATION_EXPIRY_HOURS} hour(s)."
    )
    write_email(user.email, subject, body)

def issue_password_reset(db: Session, user: models.User) -> None:
    db.query(models.PasswordResetToken).filter(
        models.PasswordResetToken.user_id == user.id,
        models.PasswordResetToken.used_at.is_(None),
    ).delete()
    raw_token = create_token()
    token_hash = hash_token(raw_token)
    expires_at = datetime.utcnow() + timedelta(minutes=PASSWORD_RESET_EXPIRY_MINUTES)
    db.add(
        models.PasswordResetToken(
            user_id=user.id,
            token_hash=token_hash,
            expires_at=expires_at,
        )
    )
    db.commit()
    reset_link = f"{WEB_APP_URL}/reset-password?token={raw_token}"
    subject = "Reset your Erasmus password"
    body = (
        "You requested a password reset.\n\n"
        f"{reset_link}\n\n"
        f"This link expires in {PASSWORD_RESET_EXPIRY_MINUTES} minute(s)."
    )
    write_email(user.email, subject, body)

def consume_email_verification(db: Session, token: str) -> models.User:
    token_hash = hash_token(token)
    record = (
        db.query(models.EmailVerificationToken)
        .filter(models.EmailVerificationToken.token_hash == token_hash)
        .first()
    )
    if not record or record.used_at is not None:
        raise HTTPException(status_code=400, detail="Invalid or expired token")
    if record.expires_at < datetime.utcnow():
        raise HTTPException(status_code=400, detail="Token expired")
    user = db.query(models.User).filter(models.User.id == record.user_id).first()
    if not user:
        raise HTTPException(status_code=400, detail="Invalid token")
    user.is_verified = True
    record.used_at = datetime.utcnow()
    db.commit()
    return user

def consume_password_reset(db: Session, token: str, new_password: str) -> None:
    token_hash = hash_token(token)
    record = (
        db.query(models.PasswordResetToken)
        .filter(models.PasswordResetToken.token_hash == token_hash)
        .first()
    )
    if not record or record.used_at is not None:
        raise HTTPException(status_code=400, detail="Invalid or expired token")
    if record.expires_at < datetime.utcnow():
        raise HTTPException(status_code=400, detail="Token expired")
    user = db.query(models.User).filter(models.User.id == record.user_id).first()
    if not user:
        raise HTTPException(status_code=400, detail="Invalid token")
    user.password_hash = hash_password(new_password)
    record.used_at = datetime.utcnow()
    db.query(models.PasswordResetToken).filter(
        models.PasswordResetToken.user_id == user.id
    ).delete()
    db.commit()
def ensure_login_timestamp(db: Session, user: models.User) -> None:
    if user.last_login_at is not None:
        return
    user.last_login_at = datetime.utcnow()
    db.commit()

def send_push_notification(db: Session, recipient_id: int, subject: str, body: str) -> None:
    tokens = db.query(models.PushToken).filter(models.PushToken.user_id == recipient_id).all()
    for push_token in tokens:
        payload = json.dumps({
            "to": push_token.token,
            "title": subject,
            "body": body,
            "sound": "default",
            "data": {"inbox": True},
        }).encode("utf-8")
        request = UrlRequest(
            "https://exp.host/--/api/v2/push/send",
            data=payload,
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        try:
            with urlopen(request, timeout=5) as response:
                if response.status >= 400:
                    raise RuntimeError(f"Expo push request failed with status {response.status}")
        except OSError as error:
            print(f"Push delivery failed for token {push_token.id}: {error}")

def add_inbox_message(db: Session, recipient_id: int, subject: str, body: str) -> None:
    message = models.InboxMessage(
        recipient_user_id=recipient_id,
        subject=subject,
        body=body,
    )
    db.add(message)
    db.flush()
    send_push_notification(db, recipient_id, subject, body)

def build_cta(label: str, path: str) -> str:
    return f"{label}: {WEB_APP_URL}{path}"

def safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("_")
    return cleaned or "document"

def get_db_path() -> Path:
    url = database.DATABASE_URL
    if not url.startswith("sqlite:///"):
        raise HTTPException(status_code=500, detail="Unsupported database")
    raw_path = url.replace("sqlite:///", "", 1)
    db_path = Path(raw_path)
    if not db_path.is_absolute():
        db_path = Path(__file__).resolve().parent / db_path
    return db_path

def safe_extract_zip(archive: zipfile.ZipFile, destination: Path) -> None:
    destination = destination.resolve()
    for member in archive.infolist():
        member_path = (destination / member.filename).resolve()
        if not str(member_path).startswith(str(destination)):
            raise HTTPException(status_code=400, detail="Invalid archive path")
        if member.is_dir():
            member_path.mkdir(parents=True, exist_ok=True)
        else:
            member_path.parent.mkdir(parents=True, exist_ok=True)
            with archive.open(member, "r") as source, open(member_path, "wb") as target:
                shutil.copyfileobj(source, target)

def normalize_required_docs(value: Optional[list[str]]) -> list[str]:
    if not value:
        return list(DOC_TYPES)
    normalized = []
    invalid = []
    for item in value:
        key = item.strip().lower()
        if key in DOC_TYPES:
            if key not in normalized:
                normalized.append(key)
        else:
            invalid.append(item)
    if invalid:
        invalid_list = ", ".join(str(item) for item in invalid)
        raise HTTPException(status_code=400, detail=f"Invalid document key(s): {invalid_list}")
    return normalized or list(DOC_TYPES)

def serialize_required_docs(value: list[str]) -> str:
    return ",".join(value)

def deserialize_required_docs(value: Optional[str]) -> list[str]:
    if not value:
        return list(DOC_TYPES)
    items = [item.strip() for item in value.split(",") if item.strip()]
    normalized = [item for item in items if item in DOC_TYPES]
    return normalized or list(DOC_TYPES)

def build_call_out(call: models.ErasmusCall) -> ErasmusCallOut:
    return ErasmusCallOut(
        id=call.id,
        title=call.title,
        country=call.country,
        description=call.description,
        deadline=call.deadline,
        link=call.link,
        erasmus_email=call.erasmus_email,
        min_gpa=call.min_gpa or 0,
        required_docs=deserialize_required_docs(call.required_docs),
    )

def pdf_safe(value: str) -> str:
    return (
        unicodedata.normalize("NFKD", value)
        .encode("ascii", "ignore")
        .decode("ascii")
    )

def validate_full_name(value: str) -> None:
    if not NAME_PATTERN.match(value):
        raise HTTPException(
            status_code=400,
            detail="Name can only contain letters, spaces, dots, and hyphens.",
        )
    parts = [p for p in value.strip().split() if len(p) > 0]
    if len(parts) < 2 or any(len(re.sub(r"[^A-Za-zÀ-ÖØ-öø-ÿĀ-ž]", "", p)) < 2 for p in parts):
        raise HTTPException(
            status_code=400,
            detail="Please provide your full name (first and last name), each at least 2 letters.",
        )

def validate_phone(value: str) -> None:
    if not PHONE_PATTERN.match(value):
        raise HTTPException(
            status_code=400,
            detail="Phone must contain only digits and an optional leading '+'.",
        )

def validate_gpa(gpa: float) -> None:
    if gpa < 0 or gpa > 10:
        raise HTTPException(status_code=400, detail="GPA must be between 0 and 10")

def validate_call_link(value: str) -> None:
    parsed = urlparse(value.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise HTTPException(status_code=400, detail="Link must be a valid http(s) URL")

def validate_university_email(email: str) -> None:
    if not email.endswith(ALLOWED_EMAIL_DOMAINS):
        raise HTTPException(
            status_code=400,
            detail="Use a university email address ending in @ibu.edu.ba or @stu.ibu.edu.ba",
        )

def get_active_login_lock(db: Session, email: str) -> Optional[models.LoginLock]:
    lock = db.query(models.LoginLock).filter(models.LoginLock.email == email).first()
    if lock and lock.locked_until > datetime.utcnow():
        return lock
    if lock:
        db.delete(lock)
        db.commit()
    return None

def register_failed_login(db: Session, email: str) -> None:
    db.add(models.LoginAttempt(email=email))
    db.commit()
    window_start = datetime.utcnow() - timedelta(minutes=LOGIN_LOCK_MINUTES)
    attempts = (
        db.query(models.LoginAttempt)
        .filter(models.LoginAttempt.email == email, models.LoginAttempt.attempted_at >= window_start)
        .count()
    )
    if attempts >= LOGIN_MAX_ATTEMPTS:
        lock = db.query(models.LoginLock).filter(models.LoginLock.email == email).first()
        locked_until = datetime.utcnow() + timedelta(minutes=LOGIN_LOCK_MINUTES)
        if lock:
            lock.locked_until = locked_until
        else:
            db.add(models.LoginLock(email=email, locked_until=locked_until))
        db.commit()

def clear_login_attempts(db: Session, email: str) -> None:
    db.query(models.LoginAttempt).filter(models.LoginAttempt.email == email).delete()
    lock = db.query(models.LoginLock).filter(models.LoginLock.email == email).first()
    if lock:
        db.delete(lock)
    db.commit()

def log_login_audit(
    db: Session,
    email: str,
    success: bool,
    reason: Optional[str],
    request: Request,
) -> None:
    ip_address = request.client.host if request.client else None
    user_agent = request.headers.get("user-agent")
    db.add(
        models.LoginAuditLog(
            email=email,
            success=success,
            ip_address=ip_address,
            user_agent=user_agent,
            reason=reason,
        )
    )
    db.commit()

def add_status_history(
    db: Session,
    application_id: int,
    status: str,
    actor_email: str,
    actor_role: str,
    reason: Optional[str] = None,
) -> None:
    db.add(
        models.ApplicationStatusHistory(
            application_id=application_id,
            status=status,
            reason=reason,
            changed_by_email=actor_email,
            changed_by_role=actor_role,
        )
    )

def review_application_record(
    db: Session,
    application: models.Application,
    normalized_status: str,
    reason: Optional[str],
    reviewer: models.User,
) -> models.Application:
    call = application.call
    if normalized_status == "approved" and not call.erasmus_email:
        raise HTTPException(
            status_code=400,
            detail="Call has no Erasmus contact email",
        )
    stored_reason = reason if normalized_status == "rejected" else None

    application.status = normalized_status
    application.reviewed_at = datetime.utcnow()
    application.reviewed_by_email = reviewer.email
    application.decision_reason = stored_reason
    add_status_history(
        db,
        application.id,
        normalized_status,
        reviewer.email,
        reviewer.role,
        reason=stored_reason,
    )
    audit_log = models.AuditLog(
        action="application_reviewed",
        application_id=application.id,
        call_title=call.title,
        student_email=application.user.email,
        reviewer_user_id=reviewer.id,
        reviewer_email=reviewer.email,
        status=normalized_status,
        reason=reason,
    )
    db.add(audit_log)
    db.commit()
    db.refresh(application)

    student_subject = f"Application {normalized_status}: {call.title}"
    student_body = (
        f"Your application for '{call.title}' has been {normalized_status}.\n"
        f"Deadline: {call.deadline}\n"
        f"Call link: {call.link}\n"
    )
    if stored_reason:
        student_body += f"Reason: {stored_reason}\n"
    student_body += f"\n{build_cta('Open dashboard', '/student')}\n"
    write_email(application.user.email, student_subject, student_body)
    if application.user.last_login_at is not None:
        inbox_subject = f"Application {normalized_status}: {call.title}"
        inbox_body = student_body
        add_inbox_message(db, application.user.id, inbox_subject, inbox_body)
        db.commit()

    if normalized_status == "approved" and call.erasmus_email:
        erasmus_subject = f"Approved application: {call.title}"
        erasmus_body = (
            f"Approved student application for '{call.title}'.\n\n"
            f"Student: {application.full_name}\n"
            f"Email: {application.user.email}\n"
            f"Department: {application.department}\n"
            f"GPA: {application.gpa}\n"
            f"Phone: {application.phone}\n"
            f"Notes: {application.notes or '-'}\n"
            f"Applied at: {application.created_at}\n"
        )
        write_email(call.erasmus_email, erasmus_subject, erasmus_body)

    return application

def parse_deadline_date(value: str) -> Optional[datetime]:
    try:
        return datetime.strptime(value, "%Y-%m-%d")
    except ValueError:
        return None

def maybe_create_deadline_alerts(
    db: Session,
    current_user: models.User,
    calls: list[models.ErasmusCall],
) -> None:
    if current_user.role != "student":
        return
    today = datetime.utcnow().date()
    created = False
    for call in calls:
        deadline = parse_deadline_date(call.deadline)
        if not deadline:
            continue
        days_left = (deadline.date() - today).days
        if days_left not in ALERT_DAYS:
            continue
        existing = (
            db.query(models.CallDeadlineAlert)
            .filter(
                models.CallDeadlineAlert.user_id == current_user.id,
                models.CallDeadlineAlert.call_id == call.id,
                models.CallDeadlineAlert.days_before == days_left,
            )
            .first()
        )
        if existing:
            continue
        subject = f"Deadline reminder: {call.title}"
        body = (
            f"Only {days_left} day(s) left to apply.\n"
            f"Deadline: {call.deadline}\n"
            f"Link: {call.link}\n\n"
            f"{build_cta('Open calls', '/student')}\n"
        )
        add_inbox_message(db, current_user.id, subject, body)
        db.add(
            models.CallDeadlineAlert(
                user_id=current_user.id,
                call_id=call.id,
                days_before=days_left,
            )
        )
        created = True
    if created:
        db.commit()

def compress_image(data: bytes, ext: str) -> bytes:
    with Image.open(BytesIO(data)) as image:
        output = BytesIO()
        if ext in {".jpg", ".jpeg"}:
            rgb_image = image.convert("RGB")
            rgb_image.save(output, format="JPEG", quality=85, optimize=True)
        else:
            image.save(output, format="PNG", optimize=True)
        return output.getvalue()

def ensure_file_signature_bytes(data: bytes, ext: str) -> None:
    signature = FILE_SIGNATURES.get(ext)
    if not signature:
        return
    if not data.startswith(signature):
        raise HTTPException(status_code=400, detail="File content does not match extension")

def ensure_file_signature_stream(file: UploadFile, ext: str) -> None:
    signature = FILE_SIGNATURES.get(ext)
    if not signature:
        return
    file.file.seek(0)
    header = file.file.read(len(signature))
    file.file.seek(0)
    if not header.startswith(signature):
        raise HTTPException(status_code=400, detail="File content does not match extension")

def run_virus_scan(file_path: Path) -> None:
    if not CLAMAV_PATH:
        return
    scanner = Path(CLAMAV_PATH)
    if not scanner.exists():
        raise HTTPException(status_code=500, detail="Virus scanner not available")
    try:
        result = subprocess.run(
            [str(scanner), "--no-summary", str(file_path)],
            capture_output=True,
            text=True,
            timeout=CLAMAV_TIMEOUT,
        )
    except subprocess.TimeoutExpired:
        file_path.unlink(missing_ok=True)
        raise HTTPException(status_code=500, detail="Virus scan timed out")
    if result.returncode == 0:
        return
    if result.returncode == 1:
        file_path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail="File failed virus scan")
    raise HTTPException(status_code=500, detail="Virus scan failed")

def save_upload(file: UploadFile, doc_type: str) -> tuple[str, int, str]:
    if not file or not file.filename:
        raise HTTPException(status_code=400, detail="Missing required document")
    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail="Only PDF, DOCX, JPG, or PNG files are allowed",
        )
    content_type = (file.content_type or "").lower()
    if not content_type or content_type not in ALLOWED_MIME_TYPES.get(ext, set()):
        raise HTTPException(status_code=400, detail="Invalid file type")

    stored_name = f"{doc_type}_{uuid4().hex}{ext}"
    target = UPLOAD_DIR / stored_name
    if ext in IMAGE_EXTENSIONS:
        data = file.file.read()
        file.file.close()
        if not data:
            raise HTTPException(status_code=400, detail="Empty file uploaded")
        ensure_file_signature_bytes(data, ext)
        if len(data) > MAX_IMAGE_RAW_SIZE:
            raise HTTPException(status_code=413, detail="Image too large (max 15MB)")
        if len(data) > MAX_IMAGE_SIZE and ENABLE_IMAGE_COMPRESSION:
            try:
                data = compress_image(data, ext)
            except UnidentifiedImageError:
                raise HTTPException(status_code=400, detail="Invalid image file")
        if len(data) > MAX_IMAGE_SIZE:
            raise HTTPException(status_code=413, detail="Image too large (max 5MB)")
        target.write_bytes(data)
        run_virus_scan(target)
        return stored_name, len(data), file.filename

    ensure_file_signature_stream(file, ext)
    size = 0
    with target.open("wb") as buffer:
        while True:
            chunk = file.file.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > MAX_DOC_SIZE:
                buffer.close()
                target.unlink(missing_ok=True)
                raise HTTPException(status_code=413, detail="File too large (max 10MB)")
            buffer.write(chunk)
    file.file.close()
    if size == 0:
        target.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail="Empty file uploaded")
    run_virus_scan(target)
    return stored_name, size, file.filename

def ensure_admin(db: Session) -> None:
    existing = db.query(models.User).filter(models.User.email == ADMIN_EMAIL).first()
    if existing:
        return

    admin_user = models.User(
        email=ADMIN_EMAIL,
        password_hash=hash_password(ADMIN_PASSWORD),
        role="admin",
        is_verified=True,
    )
    db.add(admin_user)
    db.commit()

@app.on_event("startup")
def startup() -> None:
    db = database.SessionLocal()
    try:
        ensure_admin(db)
    finally:
        db.close()

def get_current_user(
    authorization: Optional[str] = Header(None),
    db: Session = Depends(get_db),
) -> models.User:
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Missing token")

    token = authorization.split(" ", 1)[1]
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        user_id = int(payload.get("sub", "0"))
    except (JWTError, ValueError):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid token")

    user = db.query(models.User).filter(models.User.id == user_id).first()
    if not user:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="User not found")
    if not user.is_verified:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Email not verified")
    ensure_login_timestamp(db, user)
    return user

def require_admin(current_user: models.User = Depends(get_current_user)) -> models.User:
    if current_user.role != "admin":
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Admin only")
    return current_user

@app.get("/")
def home():
    return {"message": "Erasmus API running"}

@app.get("/health")
def health():
    db = database.SessionLocal()
    try:
        db.execute(text("SELECT 1"))
        return {"status": "ok", "service": "erasmus-api", "database": "ok"}
    finally:
        db.close()

@app.get("/me", response_model=UserOut)
def get_me(current_user: models.User = Depends(get_current_user)):
    return current_user

@app.post("/auth/register", response_model=MessageOut)
def register(payload: PasswordCreateRequest, db: Session = Depends(get_db)):
    email = payload.email.strip().lower()
    validate_university_email(email)

    existing = db.query(models.User).filter(models.User.email == email).first()
    if existing:
        raise HTTPException(status_code=409, detail="Email already registered")

    user = models.User(
        email=email,
        password_hash=hash_password(payload.password),
        role="student",
        is_verified=False,
    )
    db.add(user)
    db.commit()
    db.refresh(user)

    issue_email_verification(db, user)
    return {"message": "Check your email to verify your account."}

@app.post("/auth/login", response_model=AuthResponse)
def login(payload: AuthRequest, request: Request, db: Session = Depends(get_db)):
    email = payload.email.strip().lower()
    validate_university_email(email)
    lock = get_active_login_lock(db, email)
    if lock:
        log_login_audit(
            db,
            email,
            False,
            f"locked_until:{lock.locked_until.isoformat()}",
            request,
        )
        raise HTTPException(
            status_code=429,
            detail=f"Too many failed attempts. Try again after {lock.locked_until}.",
        )
    user = db.query(models.User).filter(models.User.email == email).first()
    if not user or not verify_password(payload.password, user.password_hash):
        register_failed_login(db, email)
        log_login_audit(db, email, False, "invalid_credentials", request)
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid credentials")
    if not user.is_verified:
        log_login_audit(db, email, False, "email_not_verified", request)
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Email not verified")

    if user.two_factor_enabled:
        challenge_token = create_two_factor_token(str(user.id))
        db.add(
            models.TwoFactorChallenge(
                user_id=user.id,
                token_hash=hash_token(challenge_token),
                expires_at=datetime.utcnow() + timedelta(minutes=5),
            )
        )
        db.commit()
        return {
            "access_token": "",
            "token_type": "bearer",
            "user": user,
            "requires_2fa": True,
            "two_factor_token": challenge_token,
        }

    clear_login_attempts(db, email)
    user.last_login_at = datetime.utcnow()
    db.commit()
    db.refresh(user)
    log_login_audit(db, email, True, "success", request)

    token = create_access_token(str(user.id))
    return {"access_token": token, "token_type": "bearer", "user": user}

@app.post("/auth/verify-2fa", response_model=AuthResponse)
def verify_two_factor(
    payload: TwoFactorVerifyRequest,
    request: Request,
    db: Session = Depends(get_db),
):
    try:
        token_payload = jwt.decode(
            payload.two_factor_token,
            SECRET_KEY,
            algorithms=[ALGORITHM],
        )
        if token_payload.get("purpose") != "2fa":
            raise JWTError("Invalid 2FA token")
        user_id = int(token_payload.get("sub", "0"))
    except (JWTError, ValueError):
        raise HTTPException(status_code=401, detail="Invalid or expired 2FA challenge")

    user = db.query(models.User).filter(models.User.id == user_id).first()
    if not user or not user.is_verified or not user.two_factor_enabled or not user.totp_secret:
        raise HTTPException(status_code=401, detail="2FA is not available for this account")
    challenge = (
        db.query(models.TwoFactorChallenge)
        .filter(
            models.TwoFactorChallenge.token_hash == hash_token(payload.two_factor_token),
            models.TwoFactorChallenge.user_id == user.id,
            models.TwoFactorChallenge.used_at.is_(None),
        )
        .first()
    )
    if not challenge or challenge.expires_at < datetime.utcnow():
        raise HTTPException(status_code=401, detail="Invalid or expired 2FA challenge")
    if not pyotp.TOTP(user.totp_secret).verify(payload.code.strip(), valid_window=1):
        register_failed_login(db, user.email)
        log_login_audit(db, user.email, False, "invalid_2fa_code", request)
        raise HTTPException(status_code=401, detail="Invalid 2FA code")

    challenge.used_at = datetime.utcnow()
    clear_login_attempts(db, user.email)
    user.last_login_at = datetime.utcnow()
    db.commit()
    db.refresh(user)
    log_login_audit(db, user.email, True, "success_2fa", request)
    return {
        "access_token": create_access_token(str(user.id)),
        "token_type": "bearer",
        "user": user,
    }

@app.post("/auth/2fa/setup", response_model=TwoFactorSetupOut)
def setup_two_factor(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    secret = current_user.totp_secret or pyotp.random_base32()
    current_user.totp_secret = secret
    db.commit()
    uri = pyotp.TOTP(secret).provisioning_uri(
        name=current_user.email,
        issuer_name="Erasmus Portal",
    )
    return {"secret": secret, "otpauth_uri": uri, "qr_code": build_qr_code(uri)}

@app.post("/auth/2fa/confirm", response_model=MessageOut)
def confirm_two_factor(
    payload: TwoFactorCodeRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    if not current_user.totp_secret:
        raise HTTPException(status_code=400, detail="Start 2FA setup first")
    if not pyotp.TOTP(current_user.totp_secret).verify(payload.code.strip(), valid_window=1):
        raise HTTPException(status_code=400, detail="Invalid 2FA code")
    current_user.two_factor_enabled = True
    db.commit()
    return {"message": "Two-factor authentication enabled."}

@app.post("/auth/2fa/disable", response_model=MessageOut)
def disable_two_factor(
    payload: TwoFactorCodeRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    if not current_user.two_factor_enabled or not current_user.totp_secret:
        return {"message": "Two-factor authentication is already disabled."}
    if not pyotp.TOTP(current_user.totp_secret).verify(payload.code.strip(), valid_window=1):
        raise HTTPException(status_code=400, detail="Invalid 2FA code")
    current_user.two_factor_enabled = False
    current_user.totp_secret = None
    db.commit()
    return {"message": "Two-factor authentication disabled."}

@app.post("/auth/verify-email", response_model=MessageOut)
def verify_email(payload: VerifyEmailRequest, db: Session = Depends(get_db)):
    consume_email_verification(db, payload.token.strip())
    return {"message": "Email verified. You can now log in."}

@app.post("/auth/resend-verification", response_model=MessageOut)
def resend_verification(payload: ResendVerificationRequest, db: Session = Depends(get_db)):
    email = payload.email.strip().lower()
    user = db.query(models.User).filter(models.User.email == email).first()
    if not user or user.is_verified:
        return {"message": "If the account exists, a verification email was sent."}
    issue_email_verification(db, user)
    return {"message": "Verification email sent."}

@app.post("/auth/request-password-reset", response_model=MessageOut)
def request_password_reset(payload: PasswordResetRequest, db: Session = Depends(get_db)):
    email = payload.email.strip().lower()
    user = db.query(models.User).filter(models.User.email == email).first()
    if user:
        issue_password_reset(db, user)
    return {"message": "If the account exists, a reset email was sent."}

@app.post("/auth/reset-password", response_model=MessageOut)
def reset_password(payload: PasswordResetConfirm, db: Session = Depends(get_db)):
    consume_password_reset(db, payload.token.strip(), payload.new_password)
    return {"message": "Password updated. You can now log in."}

@app.post("/auth/change-password", response_model=MessageOut)
def change_password(
    payload: ChangePasswordRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    if not verify_password(payload.current_password, current_user.password_hash):
        raise HTTPException(status_code=400, detail="Current password is incorrect")
    current_user.password_hash = hash_password(payload.new_password)
    db.commit()
    return {"message": "Password updated successfully."}

@app.get("/calls", response_model=list[ErasmusCallOut])
def get_calls(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    calls = db.query(models.ErasmusCall).order_by(models.ErasmusCall.id.desc()).all()
    maybe_create_deadline_alerts(db, current_user, calls)
    return [build_call_out(call) for call in calls]

@app.get("/calls/{call_id}", response_model=ErasmusCallOut)
def get_call(
    call_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    call = db.query(models.ErasmusCall).filter(models.ErasmusCall.id == call_id).first()
    if not call:
        raise HTTPException(status_code=404, detail="Call not found")
    return build_call_out(call)

@app.post("/calls", response_model=ErasmusCallOut)
def add_call(
    call: ErasmusCallCreate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    if not call.erasmus_email:
        raise HTTPException(status_code=400, detail="Erasmus contact email required")
    validate_call_link(call.link)
    required_docs = normalize_required_docs(call.required_docs)
    new_call = models.ErasmusCall(
        title=call.title,
        country=call.country,
        description=call.description,
        deadline=call.deadline,
        link=call.link,
        erasmus_email=call.erasmus_email,
        min_gpa=call.min_gpa,
        required_docs=serialize_required_docs(required_docs),
    )
    db.add(new_call)
    db.commit()
    db.refresh(new_call)

    notification = models.Notification(
        title=f"Novi Erasmus poziv: {new_call.title}",
        body=(
            f"{new_call.title} ({new_call.country}) - rok: {new_call.deadline}\n"
            f"{build_cta('Open calls', '/student')}"
        ),
    )
    db.add(notification)
    db.commit()

    students = (
        db.query(models.User)
        .filter(models.User.role == "student", models.User.last_login_at.isnot(None))
        .all()
    )
    subject = f"New Erasmus call: {new_call.title}"
    body = (
        f"A new Erasmus call is available.\n\n"
        f"Title: {new_call.title}\n"
        f"Country: {new_call.country}\n"
        f"Deadline: {new_call.deadline}\n"
        f"Link: {new_call.link}\n\n"
        f"{build_cta('Open calls', '/student')}\n"
    )
    for student in students:
        write_email(student.email, subject, body)
        inbox_subject = f"New Erasmus call: {new_call.title}"
        inbox_body = (
            f"A new Erasmus call has been added.\n\n"
            f"Title: {new_call.title}\n"
            f"Country: {new_call.country}\n"
            f"Deadline: {new_call.deadline}\n"
            f"Link: {new_call.link}\n\n"
            f"{build_cta('Open calls', '/student')}\n"
        )
        add_inbox_message(db, student.id, inbox_subject, inbox_body)

    if students:
        db.commit()

    return build_call_out(new_call)

@app.put("/calls/{call_id}", response_model=ErasmusCallOut)
def update_call(
    call_id: int,
    update: ErasmusCallUpdate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    call = db.query(models.ErasmusCall).filter(models.ErasmusCall.id == call_id).first()
    if not call:
        raise HTTPException(status_code=404, detail="Call not found")

    data = update.dict(exclude_unset=True)
    if "link" in data:
        validate_call_link(str(data["link"]))
    if "required_docs" in data:
        data["required_docs"] = serialize_required_docs(
            normalize_required_docs(data["required_docs"])
        )
    for key, value in data.items():
        setattr(call, key, value)
    db.commit()
    db.refresh(call)
    return build_call_out(call)

@app.delete("/calls/{call_id}", status_code=204)
def delete_call(
    call_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    call = db.query(models.ErasmusCall).filter(models.ErasmusCall.id == call_id).first()
    if not call:
        raise HTTPException(status_code=404, detail="Call not found")
    db.delete(call)
    db.commit()
    return None

@app.post("/calls/{call_id}/apply", response_model=ApplicationOut)
def apply_to_call(
    call_id: int,
    full_name: str = Form(...),
    phone: str = Form(...),
    department: str = Form(...),
    gpa: float = Form(...),
    notes: Optional[str] = Form(None),
    cv: Optional[UploadFile] = File(None),
    transcript: Optional[UploadFile] = File(None),
    motivation_letter: Optional[UploadFile] = File(None),
    proof_of_language: Optional[UploadFile] = File(None),
    recommendation_letter: Optional[UploadFile] = File(None),
    learning_agreement: Optional[UploadFile] = File(None),
    id_copy: Optional[UploadFile] = File(None),
    application_form: Optional[UploadFile] = File(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    if current_user.role != "student":
        raise HTTPException(status_code=403, detail="Students only")

    call = db.query(models.ErasmusCall).filter(models.ErasmusCall.id == call_id).first()
    if not call:
        raise HTTPException(status_code=404, detail="Call not found")

    existing = db.query(models.Application).filter(
        models.Application.call_id == call_id,
        models.Application.user_id == current_user.id,
    ).first()
    if existing:
        raise HTTPException(status_code=409, detail="Already applied")

    trimmed_full_name = full_name.strip()
    trimmed_phone = phone.strip()
    trimmed_department = department.strip()
    if not trimmed_full_name or not trimmed_phone or not trimmed_department:
        raise HTTPException(status_code=400, detail="Missing required fields")
    validate_full_name(trimmed_full_name)
    validate_phone(trimmed_phone)
    validate_gpa(gpa)
    min_gpa = call.min_gpa or 0
    if gpa < min_gpa:
        raise HTTPException(
            status_code=400,
            detail=f"GPA must be at least {min_gpa}",
        )

    required_docs = deserialize_required_docs(call.required_docs)
    doc_uploads = {
        "cv": cv,
        "transcript": transcript,
        "motivation": motivation_letter,
        "language_proof": proof_of_language,
        "recommendation": recommendation_letter,
        "learning_agreement": learning_agreement,
        "id_copy": id_copy,
        "application_form": application_form,
    }
    missing = [doc for doc in required_docs if not doc_uploads.get(doc)]
    if missing:
        labels = ", ".join(DOC_LABELS.get(doc, doc) for doc in missing)
        raise HTTPException(
            status_code=400, detail=f"Missing required documents: {labels}"
        )

    application = models.Application(
        call_id=call_id,
        user_id=current_user.id,
        full_name=trimmed_full_name,
        phone=trimmed_phone,
        department=trimmed_department,
        gpa=gpa,
        notes=notes.strip() if notes else None,
    )
    db.add(application)
    db.commit()
    db.refresh(application)

    documents = []
    for doc_type, upload in doc_uploads.items():
        if not upload:
            continue
        stored_name, size, original_name = save_upload(upload, doc_type)
        documents.append(
            models.ApplicationDocument(
                application_id=application.id,
                doc_type=doc_type,
                original_name=original_name,
                stored_name=stored_name,
                size=size,
            )
        )
    if documents:
        db.add_all(documents)
    add_status_history(
        db,
        application.id,
        "pending",
        current_user.email,
        current_user.role,
    )
    db.commit()

    admins = (
        db.query(models.User)
        .filter(models.User.role == "admin", models.User.last_login_at.isnot(None))
        .all()
    )
    if admins:
        inbox_subject = f"New application: {call.title}"
        inbox_body = (
            f"A student applied for '{call.title}'.\n\n"
            f"Student: {current_user.email}\n"
            f"Name: {application.full_name}\n"
            f"Department: {application.department}\n"
            f"GPA: {application.gpa}\n"
            f"Phone: {application.phone}\n"
            f"Notes: {application.notes or '-'}\n\n"
            f"{build_cta('Review application', f'/applications/{application.id}')}\n"
        )
        for admin in admins:
            add_inbox_message(db, admin.id, inbox_subject, inbox_body)
            write_email(admin.email, inbox_subject, inbox_body)
        db.commit()

    return ApplicationOut(
        id=application.id,
        call_id=call.id,
        call_title=call.title,
        user_email=current_user.email,
        full_name=application.full_name,
        phone=application.phone,
        department=application.department,
        gpa=application.gpa,
        notes=application.notes,
        status=application.status,
        created_at=application.created_at,
        reviewed_at=application.reviewed_at,
        reviewed_by_email=application.reviewed_by_email,
        decision_reason=application.decision_reason,
    )

@app.get("/applications", response_model=list[ApplicationOut])
def get_applications(
    response: Response,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    response.headers["Cache-Control"] = "no-store"
    rows = (
        db.query(models.Application, models.ErasmusCall, models.User)
        .join(models.ErasmusCall, models.Application.call_id == models.ErasmusCall.id)
        .join(models.User, models.Application.user_id == models.User.id)
        .order_by(models.Application.id.desc())
        .all()
    )

    results = []
    for application, call, user in rows:
        results.append(
            ApplicationOut(
                id=application.id,
                call_id=call.id,
                call_title=call.title,
                user_email=user.email,
                full_name=application.full_name,
                phone=application.phone,
                department=application.department,
                gpa=application.gpa,
                notes=application.notes,
                status=application.status,
                created_at=application.created_at,
                reviewed_at=application.reviewed_at,
                reviewed_by_email=application.reviewed_by_email,
                decision_reason=application.decision_reason,
            )
        )
    return results

@app.get("/applications/export/csv")
def export_applications_csv(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    rows = (
        db.query(models.Application, models.ErasmusCall, models.User)
        .join(models.ErasmusCall, models.Application.call_id == models.ErasmusCall.id)
        .join(models.User, models.Application.user_id == models.User.id)
        .order_by(models.Application.id.desc())
        .all()
    )
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "id",
            "call_id",
            "call_title",
            "call_country",
            "student_email",
            "full_name",
            "department",
            "phone",
            "gpa",
            "status",
            "created_at",
            "reviewed_at",
            "reviewed_by_email",
            "decision_reason",
        ]
    )
    for application, call, user in rows:
        writer.writerow(
            [
                application.id,
                call.id,
                call.title,
                call.country,
                user.email,
                application.full_name,
                application.department,
                application.phone,
                application.gpa,
                application.status,
                application.created_at,
                application.reviewed_at,
                application.reviewed_by_email,
                application.decision_reason,
            ]
        )
    headers = {"Content-Disposition": "attachment; filename=applications.csv"}
    return Response(output.getvalue(), media_type="text/csv", headers=headers)

@app.get("/applications/export/xlsx")
def export_applications_xlsx(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    rows = (
        db.query(models.Application, models.ErasmusCall, models.User)
        .join(models.ErasmusCall, models.Application.call_id == models.ErasmusCall.id)
        .join(models.User, models.Application.user_id == models.User.id)
        .order_by(models.Application.id.desc())
        .all()
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Applications"
    sheet.append(
        [
            "id",
            "call_id",
            "call_title",
            "call_country",
            "student_email",
            "full_name",
            "department",
            "phone",
            "gpa",
            "status",
            "created_at",
            "reviewed_at",
            "reviewed_by_email",
            "decision_reason",
        ]
    )
    for application, call, user in rows:
        sheet.append(
            [
                application.id,
                call.id,
                call.title,
                call.country,
                user.email,
                application.full_name,
                application.department,
                application.phone,
                application.gpa,
                application.status,
                application.created_at.isoformat(),
                application.reviewed_at.isoformat() if application.reviewed_at else "",
                application.reviewed_by_email or "",
                application.decision_reason or "",
            ]
        )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    headers = {"Content-Disposition": "attachment; filename=applications.xlsx"}
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )

@app.get("/applications/export/pdf")
def export_applications_pdf(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    rows = (
        db.query(models.Application, models.ErasmusCall, models.User)
        .join(models.ErasmusCall, models.Application.call_id == models.ErasmusCall.id)
        .join(models.User, models.Application.user_id == models.User.id)
        .order_by(models.Application.id.desc())
        .all()
    )
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 8, "Applications report", ln=True)
    pdf.set_font("Helvetica", size=10)
    if not rows:
        pdf.multi_cell(0, 6, "No applications found.")
    for application, call, user in rows:
        pdf.multi_cell(
            0,
            6,
            pdf_safe(
                f"ID: {application.id} | Call: {call.title} | Country: {call.country}"
            ),
        )
        pdf.multi_cell(
            0,
            6,
            pdf_safe(
                "Student: "
                f"{user.email} | {application.full_name} | Dept: {application.department}"
            ),
        )
        pdf.multi_cell(
            0,
            6,
            pdf_safe(
                "Status: "
                f"{application.status} | GPA: {application.gpa} | "
                f"Reviewed: {application.reviewed_at or '-'}"
            ),
        )
        if application.reviewed_by_email:
            pdf.multi_cell(
                0,
                6,
                pdf_safe(f"Reviewed by: {application.reviewed_by_email}"),
            )
        if application.decision_reason:
            pdf.multi_cell(
                0,
                6,
                pdf_safe(f"Reason: {application.decision_reason}"),
            )
        pdf.ln(1)
    data = pdf.output(dest="S").encode("latin-1")
    headers = {"Content-Disposition": "attachment; filename=applications.pdf"}
    return Response(data, media_type="application/pdf", headers=headers)

@app.get("/admin/backup")
def download_backup(
    current_user: models.User = Depends(require_admin),
):
    db_path = get_db_path()
    if not db_path.exists():
        raise HTTPException(status_code=404, detail="Database not found")
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.write(db_path, "database/erasmus.db")
        if UPLOAD_DIR.exists():
            for file_path in UPLOAD_DIR.rglob("*"):
                if file_path.is_file():
                    archive.write(
                        file_path,
                        f"uploads/{file_path.relative_to(UPLOAD_DIR)}",
                    )
    buffer.seek(0)
    timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    headers = {
        "Content-Disposition": f"attachment; filename=backup_{timestamp}.zip"
    }
    return StreamingResponse(buffer, media_type="application/zip", headers=headers)

@app.post("/admin/restore", status_code=204)
def restore_backup(
    file: UploadFile = File(...),
    current_user: models.User = Depends(require_admin),
):
    if not file.filename or not file.filename.lower().endswith(".zip"):
        raise HTTPException(status_code=400, detail="ZIP file required")
    with tempfile.TemporaryDirectory() as temp_dir:
        archive_path = Path(temp_dir) / "backup.zip"
        with open(archive_path, "wb") as target:
            shutil.copyfileobj(file.file, target)
        try:
            with zipfile.ZipFile(archive_path, "r") as archive:
                safe_extract_zip(archive, Path(temp_dir))
        except zipfile.BadZipFile:
            raise HTTPException(status_code=400, detail="Invalid ZIP file")
        extracted_db = Path(temp_dir) / "database" / "erasmus.db"
        if not extracted_db.exists():
            extracted_db = Path(temp_dir) / "erasmus.db"
        if not extracted_db.exists():
            raise HTTPException(status_code=400, detail="Database file missing")
        extracted_uploads = Path(temp_dir) / "uploads"
        db_path = get_db_path()
        database.engine.dispose()
        if db_path.exists():
            db_path.unlink()
        db_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(extracted_db, db_path)
        if UPLOAD_DIR.exists():
            shutil.rmtree(UPLOAD_DIR)
        if extracted_uploads.exists():
            shutil.copytree(extracted_uploads, UPLOAD_DIR)
        else:
            UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    database.ensure_schema()
    return None

@app.get("/applications/{application_id:int}", response_model=ApplicationDetailOut)
def get_application_detail(
    application_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    application = (
        db.query(models.Application)
        .filter(models.Application.id == application_id)
        .first()
    )
    if not application:
        raise HTTPException(status_code=404, detail="Application not found")

    call = application.call
    user = application.user
    documents = [ApplicationDocumentOut.from_orm(doc) for doc in application.documents]
    return ApplicationDetailOut(
        id=application.id,
        call_id=call.id,
        call_title=call.title,
        call_country=call.country,
        call_deadline=call.deadline,
        call_description=call.description,
        call_link=call.link,
        call_erasmus_email=call.erasmus_email,
        user_email=user.email,
        full_name=application.full_name,
        phone=application.phone,
        department=application.department,
        gpa=application.gpa,
        notes=application.notes,
        status=application.status,
        created_at=application.created_at,
        reviewed_at=application.reviewed_at,
        reviewed_by_email=application.reviewed_by_email,
        decision_reason=application.decision_reason,
        documents=documents,
    )

@app.get(
    "/applications/{application_id:int}/history",
    response_model=list[ApplicationStatusHistoryOut],
)
def get_application_history(
    application_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    application = (
        db.query(models.Application)
        .filter(models.Application.id == application_id)
        .first()
    )
    if not application:
        raise HTTPException(status_code=404, detail="Application not found")
    if current_user.role != "admin" and application.user_id != current_user.id:
        raise HTTPException(status_code=403, detail="Forbidden")
    return (
        db.query(models.ApplicationStatusHistory)
        .filter(models.ApplicationStatusHistory.application_id == application_id)
        .order_by(models.ApplicationStatusHistory.id.asc())
        .all()
    )

@app.get("/applications/{application_id:int}/documents.zip")
def download_application_zip(
    application_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    application = (
        db.query(models.Application)
        .filter(models.Application.id == application_id)
        .first()
    )
    if not application:
        raise HTTPException(status_code=404, detail="Application not found")
    if not application.documents:
        raise HTTPException(status_code=404, detail="No documents available")

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for doc in application.documents:
            file_path = UPLOAD_DIR / doc.stored_name
            if not file_path.exists():
                raise HTTPException(status_code=404, detail="File not found")
            filename = safe_filename(doc.original_name)
            archive.write(file_path, f"{doc.doc_type}_{filename}")
    buffer.seek(0)
    headers = {
        "Content-Disposition": f"attachment; filename=application_{application_id}_documents.zip"
    }
    return StreamingResponse(buffer, media_type="application/zip", headers=headers)

@app.post("/applications/{application_id:int}/review", response_model=ApplicationOut)
def review_application(
    application_id: int,
    payload: ApplicationReviewRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    normalized_status = payload.status.strip().lower()
    if normalized_status not in {"approved", "rejected"}:
        raise HTTPException(status_code=400, detail="Invalid status")
    reason = payload.reason.strip() if payload.reason else None
    if normalized_status == "rejected" and not reason:
        raise HTTPException(status_code=400, detail="Rejection reason required")

    application = (
        db.query(models.Application)
        .filter(models.Application.id == application_id)
        .first()
    )
    if not application:
        raise HTTPException(status_code=404, detail="Application not found")
    if application.status != "pending":
        raise HTTPException(status_code=409, detail="Application already reviewed")

    application = review_application_record(
        db,
        application,
        normalized_status,
        reason,
        current_user,
    )

    return ApplicationOut(
        id=application.id,
        call_id=application.call.id,
        call_title=application.call.title,
        user_email=application.user.email,
        full_name=application.full_name,
        phone=application.phone,
        department=application.department,
        gpa=application.gpa,
        notes=application.notes,
        status=application.status,
        created_at=application.created_at,
        reviewed_at=application.reviewed_at,
        reviewed_by_email=application.reviewed_by_email,
        decision_reason=application.decision_reason,
    )

@app.post("/applications/bulk-review", response_model=BulkReviewResult)
def bulk_review_applications(
    payload: BulkReviewRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    normalized_status = payload.status.strip().lower()
    if normalized_status not in {"approved", "rejected"}:
        raise HTTPException(status_code=400, detail="Invalid status")
    reason = payload.reason.strip() if payload.reason else None
    if normalized_status == "rejected" and not reason:
        raise HTTPException(status_code=400, detail="Rejection reason required")
    if not payload.ids:
        raise HTTPException(status_code=400, detail="No applications selected")

    unique_ids = sorted(set(payload.ids))
    applications = (
        db.query(models.Application)
        .filter(models.Application.id.in_(unique_ids))
        .all()
    )
    by_id = {app.id: app for app in applications}
    updated: list[int] = []
    skipped: list[dict] = []
    for app_id in unique_ids:
        application = by_id.get(app_id)
        if not application:
            skipped.append({"id": app_id, "reason": "Application not found"})
            continue
        if application.status != "pending":
            skipped.append({"id": app_id, "reason": "Already reviewed"})
            continue
        try:
            review_application_record(
                db,
                application,
                normalized_status,
                reason,
                current_user,
            )
            updated.append(app_id)
        except HTTPException as exc:
            skipped.append({"id": app_id, "reason": str(exc.detail)})
    return {"updated": updated, "skipped": skipped}

@app.get("/documents/{document_id}")
def download_document(
    document_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    document = (
        db.query(models.ApplicationDocument)
        .filter(models.ApplicationDocument.id == document_id)
        .first()
    )
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")

    application = document.application
    if current_user.role != "admin" and application.user_id != current_user.id:
        raise HTTPException(status_code=403, detail="Forbidden")

    file_path = UPLOAD_DIR / document.stored_name
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")

    media_type = mimetypes.guess_type(document.original_name)[0] or "application/octet-stream"
    return FileResponse(
        file_path,
        media_type=media_type,
        headers={"Content-Disposition": f'inline; filename="{safe_filename(document.original_name)}"'},
    )

@app.post("/applications/{application_id:int}/withdraw", response_model=ApplicationOut)
def withdraw_application(
    application_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    if current_user.role != "student":
        raise HTTPException(status_code=403, detail="Students only")

    application = (
        db.query(models.Application)
        .filter(
            models.Application.id == application_id,
            models.Application.user_id == current_user.id,
        )
        .first()
    )
    if not application:
        raise HTTPException(status_code=404, detail="Application not found")
    if application.status != "pending":
        raise HTTPException(status_code=409, detail="Application cannot be withdrawn")

    application.status = "withdrawn"
    application.reviewed_at = datetime.utcnow()
    application.reviewed_by_email = current_user.email
    add_status_history(
        db,
        application.id,
        "withdrawn",
        current_user.email,
        current_user.role,
        reason="Student withdrew",
    )
    db.commit()
    db.refresh(application)

    admins = (
        db.query(models.User)
        .filter(models.User.role == "admin", models.User.last_login_at.isnot(None))
        .all()
    )
    if admins:
        call = application.call
        inbox_subject = f"Application withdrawn: {call.title}"
        inbox_body = (
            f"Student {current_user.email} withdrew their application for '{call.title}'.\n\n"
            f"Name: {application.full_name}\n"
            f"Department: {application.department}\n"
            f"Phone: {application.phone}\n"
            f"Notes: {application.notes or '-'}\n\n"
            f"{build_cta('View application', f'/applications/{application.id}')}\n"
        )
        for admin in admins:
            add_inbox_message(db, admin.id, inbox_subject, inbox_body)
            write_email(admin.email, inbox_subject, inbox_body)
        db.commit()

    return ApplicationOut(
        id=application.id,
        call_id=application.call.id,
        call_title=application.call.title,
        user_email=application.user.email,
        full_name=application.full_name,
        phone=application.phone,
        department=application.department,
        gpa=application.gpa,
        notes=application.notes,
        status=application.status,
        created_at=application.created_at,
        reviewed_at=application.reviewed_at,
        reviewed_by_email=application.reviewed_by_email,
        decision_reason=application.decision_reason,
    )

@app.get("/applications/me", response_model=list[ApplicationOut])
def get_my_applications(
    response: Response,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    response.headers["Cache-Control"] = "no-store"
    rows = (
        db.query(models.Application, models.ErasmusCall)
        .join(models.ErasmusCall, models.Application.call_id == models.ErasmusCall.id)
        .filter(models.Application.user_id == current_user.id)
        .order_by(models.Application.id.desc())
        .all()
    )

    results = []
    for application, call in rows:
        results.append(
            ApplicationOut(
                id=application.id,
                call_id=call.id,
                call_title=call.title,
                user_email=current_user.email,
                full_name=application.full_name,
                phone=application.phone,
                department=application.department,
                gpa=application.gpa,
                notes=application.notes,
                status=application.status,
                created_at=application.created_at,
                reviewed_at=application.reviewed_at,
                reviewed_by_email=application.reviewed_by_email,
                decision_reason=application.decision_reason,
            )
        )
    return results

@app.get("/notifications", response_model=list[NotificationOut])
def get_notifications(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    return db.query(models.Notification).order_by(models.Notification.id.desc()).all()

@app.get("/inbox", response_model=list[InboxMessageOut])
def get_inbox(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    return (
        db.query(models.InboxMessage)
        .filter(models.InboxMessage.recipient_user_id == current_user.id)
        .order_by(models.InboxMessage.id.desc())
        .all()
    )

@app.post("/push-tokens", response_model=MessageOut)
def register_push_token(
    payload: PushTokenRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    existing = db.query(models.PushToken).filter(models.PushToken.token == payload.token).first()
    if existing:
        existing.user_id = current_user.id
        existing.platform = payload.platform
    else:
        db.add(models.PushToken(user_id=current_user.id, token=payload.token, platform=payload.platform))
    db.commit()
    return {"message": "Push token registered"}

@app.get("/inbox/unread-count", response_model=InboxUnreadCountOut)
def get_inbox_unread_count(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    count = (
        db.query(models.InboxMessage)
        .filter(
            models.InboxMessage.recipient_user_id == current_user.id,
            models.InboxMessage.read_at.is_(None),
        )
        .count()
    )
    return {"count": count}

@app.post("/inbox/{message_id:int}/read", response_model=InboxMessageOut)
def mark_inbox_read(
    message_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    message = (
        db.query(models.InboxMessage)
        .filter(
            models.InboxMessage.id == message_id,
            models.InboxMessage.recipient_user_id == current_user.id,
        )
        .first()
    )
    if not message:
        raise HTTPException(status_code=404, detail="Message not found")
    if message.read_at is None:
        message.read_at = datetime.utcnow()
        db.commit()
        db.refresh(message)
    return message

@app.post("/inbox/mark-all-read", response_model=InboxBulkActionOut)
def mark_all_inbox_read(
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    updated = (
        db.query(models.InboxMessage)
        .filter(
            models.InboxMessage.recipient_user_id == current_user.id,
            models.InboxMessage.read_at.is_(None),
        )
        .update({models.InboxMessage.read_at: datetime.utcnow()}, synchronize_session=False)
    )
    db.commit()
    return {"count": updated}

@app.post("/inbox/bulk-delete", response_model=InboxBulkActionOut)
def bulk_delete_inbox(
    payload: InboxBulkDeleteRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    if not payload.ids:
        return {"count": 0}
    deleted = (
        db.query(models.InboxMessage)
        .filter(
            models.InboxMessage.recipient_user_id == current_user.id,
            models.InboxMessage.id.in_(payload.ids),
        )
        .delete(synchronize_session=False)
    )
    db.commit()
    return {"count": deleted}

@app.delete("/inbox/{message_id:int}", status_code=204)
def delete_inbox_message(
    message_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    message = (
        db.query(models.InboxMessage)
        .filter(
            models.InboxMessage.id == message_id,
            models.InboxMessage.recipient_user_id == current_user.id,
        )
        .first()
    )
    if not message:
        raise HTTPException(status_code=404, detail="Message not found")
    db.delete(message)
    db.commit()
    return None

@app.get("/audit-log", response_model=list[AuditLogOut])
def get_audit_log(
    limit: int = 100,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    limit = max(1, min(limit, 500))
    return (
        db.query(models.AuditLog)
        .order_by(models.AuditLog.id.desc())
        .limit(limit)
        .all()
    )

@app.get("/audit-log/login", response_model=list[LoginAuditOut])
def get_login_audit_log(
    limit: int = 100,
    email: Optional[str] = None,
    success: Optional[bool] = None,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    limit = max(1, min(limit, 500))
    query = db.query(models.LoginAuditLog)
    if email:
        query = query.filter(models.LoginAuditLog.email == email.strip().lower())
    if success is not None:
        query = query.filter(models.LoginAuditLog.success == success)
    return (
        query.order_by(models.LoginAuditLog.id.desc())
        .limit(limit)
        .all()
    )

@app.delete("/notifications/{notification_id}", status_code=204)
def delete_notification(
    notification_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_admin),
):
    notification = (
        db.query(models.Notification)
        .filter(models.Notification.id == notification_id)
        .first()
    )
    if not notification:
        raise HTTPException(status_code=404, detail="Notification not found")
    db.delete(notification)
    db.commit()
    return None